import os
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

STATUS_COLORS = {
    "Resolved": "C8E6C9",
    "In Progress": "FFF9C4",
    "Monitoring": "BBDEFB",
}


def build():
    with open(os.path.join(DATA_DIR, "blocker_risk_log.csv"), newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    wb = Workbook()
    ws = wb.active
    ws.title = "Blocker & Risk Log"

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="B71C1C")
    title_font = Font(name="Arial", bold=True, size=14)
    normal_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="B0BEC5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A1"] = "Cycle-Count Process Rollout — Blocker & Risk Log"
    ws["A1"].font = title_font
    headers = list(rows[0].keys())
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))

    header_row = 3
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=i, value=h.replace("_", " "))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    r = header_row + 1
    for row in rows:
        for i, h in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=i, value=row[h])
            cell.font = normal_font
            cell.alignment = left if h in ("Description", "Impact", "Resolution_Action") else center
            cell.border = border
            if h == "Status" and row[h] in STATUS_COLORS:
                cell.fill = PatternFill("solid", fgColor=STATUS_COLORS[row[h]])
        r += 1

    widths = {"Blocker_ID": 10, "Related_Task": 12, "Description": 40, "Function_Raised_By": 18,
              "Impact": 35, "Owner": 18, "Date_Raised": 13, "Resolution_Action": 45, "Status": 13}
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 15)

    ws.freeze_panes = "A4"

    out_path = os.path.join(OUTPUT_DIR, "Blocker_Risk_Log.xlsx")
    wb.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    build()
